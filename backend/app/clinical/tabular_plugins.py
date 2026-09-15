from __future__ import annotations

import io
from collections.abc import Mapping
from typing import Any, Protocol


class TabularAdapter(Protocol):
    """Plugin contract: bytes in, column metadata and rows out."""

    available: bool
    display_name: str
    package_hint: str

    def parse(self, filename: str, content_type: str, payload: bytes) -> tuple[tuple[str, ...], tuple[Mapping[str, Any], ...]]: ...


def _module_available(module_name: str) -> bool:
    try:
        from importlib.util import find_spec

        return find_spec(module_name) is not None
    except (ImportError, ModuleNotFoundError, ValueError):
        return False


def _missing_message(display_name: str, package_hint: str) -> str:
    return f"{display_name} 适配器未安装，请安装可选依赖 {package_hint} 后重试"


class XlsxAdapter:
    display_name = "XLSX（Excel）"
    package_hint = "openpyxl>=3.1（InsightFlow 默认依赖）"
    available = _module_available("openpyxl")

    def parse(self, filename: str, content_type: str, payload: bytes):
        if not self.available:
            raise ValueError(_missing_message(self.display_name, self.package_hint))
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - guarded by available
            raise ValueError(_missing_message(self.display_name, self.package_hint)) from exc

        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            values = worksheet.iter_rows(values_only=True)
            header = next(values, None)
            if header is None:
                raise ValueError(f"{filename} 没有表头行")
            columns = tuple(str(value).strip() if value is not None else "" for value in header)
            rows = tuple({columns[index]: value for index, value in enumerate(row) if index < len(columns)} for row in values if any(value is not None for value in row))
            return columns, rows
        finally:
            workbook.close()


class ParquetAdapter:
    display_name = "Parquet"
    package_hint = "insightflow-backend[tabular]"
    available = _module_available("pyarrow")

    def parse(self, filename: str, content_type: str, payload: bytes):
        if not self.available:
            raise ValueError(_missing_message(self.display_name, self.package_hint))
        try:
            import pyarrow.parquet as parquet
        except ImportError as exc:  # pragma: no cover - guarded by available
            raise ValueError(_missing_message(self.display_name, self.package_hint)) from exc

        table = parquet.read_table(io.BytesIO(payload))
        return tuple(str(name) for name in table.column_names), tuple(table.to_pylist())


class XptAdapter:
    display_name = "XPT（SAS Transport）"
    package_hint = "insightflow-backend[tabular]"
    available = _module_available("pandas")

    def parse(self, filename: str, content_type: str, payload: bytes):
        if not self.available:
            raise ValueError(_missing_message(self.display_name, self.package_hint))
        try:
            import pandas as pd
        except ImportError as exc:  # pragma: no cover - guarded by available
            raise ValueError(_missing_message(self.display_name, self.package_hint)) from exc

        frame = pd.read_sas(io.BytesIO(payload), format="xport")
        return tuple(str(name) for name in frame.columns), tuple(frame.to_dict(orient="records"))


def optional_adapters() -> dict[str, TabularAdapter]:
    """Return lazy adapters; importing this module never imports heavy data libraries."""
    return {"xlsx": XlsxAdapter(), "parquet": ParquetAdapter(), "xpt": XptAdapter()}

